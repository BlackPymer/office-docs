#!/bin/sh
"true" '''\'
# Python-скрипт: sh-трampoline ниже запускает его нужным интерпретатором.
# Порядок: .venv рядом со скриптом -> старый пользовательский venv -> python3.
D="$(dirname "$0")"
for P in "$D/.venv/bin/python" "$HOME/.local/share/office-skills/venv/bin/python"; do
  [ -x "$P" ] && exec "$P" "$0" "$@"
done
exec python3 "$0" "$@"
''' + ''   # конкатенация -> Python не примет трampoline за docstring модуля
"""Выгрузка содержимого .xlsx: листы, диапазоны, значения и формулы.

Использование:
    xlsx_dump.py <файл.xlsx> [--sheet ИМЯ] [--formulas] [--max-rows N] [--max-cols N]

По умолчанию печатает сохранённые значения (data_only=True). С --formulas —
сами формулы. Важно: если файл только что записали openpyxl, значения
формул отсутствуют (None), пока не выполнен пересчёт — см. xlsx_recalc.py.
"""


import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ERRS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A")


def main() -> int:
    argv = sys.argv[1:]
    if not argv or argv[0].startswith("--"):
        print(__doc__)
        return 2
    path = argv[0]
    formulas = "--formulas" in argv
    sheet = None
    max_rows, max_cols = 60, 20
    if "--sheet" in argv:
        sheet = argv[argv.index("--sheet") + 1]
    if "--max-rows" in argv:
        max_rows = int(argv[argv.index("--max-rows") + 1])
    if "--max-cols" in argv:
        max_cols = int(argv[argv.index("--max-cols") + 1])

    wb = load_workbook(path, data_only=not formulas)
    print(f"# {path} — листы: {', '.join(wb.sheetnames)}")
    if sheet:
        if sheet not in wb.sheetnames:
            print(f"Нет листа «{sheet}»", file=sys.stderr)
            return 1
        targets = [sheet]
    else:
        targets = wb.sheetnames

    err_count = 0
    for name in targets:
        ws = wb[name]
        print(f"\n## {name} (диапазон: {ws.dimensions}, строк: {ws.max_row})")
        for row in ws.iter_rows(
            min_row=1, max_row=min(ws.max_row, max_rows), max_col=min(ws.max_column, max_cols)
        ):
            vals = []
            for c in row:
                v = c.value
                if v is None:
                    v = ""
                v = str(v).replace("\n", "\\n")
                if len(v) > 60:
                    v = v[:60] + "…"
                if v in ERRS:
                    err_count += 1
                    v = f"!!{v}!!"
                vals.append(v)
            while vals and vals[-1] == "":
                vals.pop()
            if not vals:
                continue
            print(f"{row[0].row:>4} | " + " | ".join(vals))
        if ws.max_row > max_rows:
            print(f"  … показано {max_rows} из {ws.max_row} строк (--max-rows)")
        if ws.max_column > max_cols:
            print(f"  … показано {max_cols} из {ws.max_column} колонок (--max-cols), "
                  f"последняя колонка: {get_column_letter(ws.max_column)}")
    if err_count:
        print(f"\nВНИМАНИЕ: найдено ячеек с ошибками: {err_count} — прогоните xlsx_recalc.py")
    return 0


if __name__ == "__main__":
    sys.exit(main())
