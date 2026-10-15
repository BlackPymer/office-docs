#!/bin/sh
"true" '''\'
# Python-скрипт: sh-трampoline ниже запускает его нужным интерпретатором.
# Порядок: .venv в корне репозитория -> .venv рядом со скриптом ->
# старый пользовательский venv -> python3.
D="$(dirname "$0")"
for P in "$D/../.venv/bin/python" "$D/.venv/bin/python" \
         "$HOME/.local/share/office-skills/venv/bin/python"; do
  [ -x "$P" ] && exec "$P" "$0" "$@"
done
exec python3 "$0" "$@"
''' + ''   # конкатенация -> Python не примет трampoline за docstring модуля
DOC = """Краткая текстовая выгрузка структуры .pptx: по слайдам, с фигурами и текстом.

Использование:
    pptx_outline.py <deck.pptx> [--max-chars 400]

Что печатает на каждый слайд: номер, layout, все фигуры (тип, имя,
плейсхолдер), их текст, заметки докладчика. Этого обычно хватает, чтобы
понять содержание и решить, что править, не открывая XML.
"""
__doc__ = DOC


import os
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def shape_text(shape, max_chars: int) -> str:
    if not shape.has_text_frame:
        return ""
    txt = "\n".join(p.text for p in shape.text_frame.paragraphs)
    txt = txt.strip()
    if len(txt) > max_chars:
        txt = txt[:max_chars] + f"… [+{len(txt) - max_chars} симв.]"
    return txt


def dump_shape(shape, indent: str, max_chars: int) -> list[str]:
    out: list[str] = []
    kind = str(shape.shape_type)
    ph = ""
    if shape.is_placeholder:
        ph = f", placeholder idx={shape.placeholder_format.idx} type={shape.placeholder_format.type}"
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        out.append(f"{indent}- GROUP «{shape.name}»{ph}")
        for sub in shape.shapes:
            out += dump_shape(sub, indent + "  ", max_chars)
        return out
    line = f"{indent}- {kind} «{shape.name}»{ph}"
    if shape.has_table:
        tbl = shape.table
        out.append(line + f" TABLE {len(tbl.rows)}x{len(tbl.columns)}")
        for r in tbl.rows:
            cells = [c.text.replace("\n", " ")[:40] for c in r.cells]
            out.append(f"{indent}    | " + " | ".join(cells))
        return out
    txt = shape_text(shape, max_chars)
    if txt:
        out.append(line)
        for ln in txt.split("\n"):
            out.append(f"{indent}    {ln}")
    else:
        out.append(line + (" (пусто)" if shape.shape_type != MSO_SHAPE_TYPE.PICTURE else ""))
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            out.append(f"{indent}    image: {shape.image.filename or shape.image.content_type}")
        except Exception:  # noqa: BLE001
            pass
    if getattr(shape, "has_chart", False) and shape.has_chart:
        ch = shape.chart
        out.append(f"{indent}    CHART type={ch.chart_type}")
    return out


def main() -> int:
    argv = sys.argv[1:]
    if not argv:
        print(__doc__)
        return 2
    path = argv[0]
    max_chars = 400
    if "--max-chars" in argv:
        max_chars = int(argv[argv.index("--max-chars") + 1])

    prs = Presentation(path)
    size = prs.slide_width, prs.slide_height
    print(f"# {path} — {len(prs.slides)} слайдов, {size[0]}x{size[1]} EMU "
          f"({size[0]/914400:.2f}x{size[1]/914400:.2f} дюймов)")
    for i, slide in enumerate(prs.slides, 1):
        print(f"\n## Слайд {i} (layout: {slide.slide_layout.name})")
        for shape in slide.shapes:
            for line in dump_shape(shape, "", max_chars):
                print(line)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                print(f"  NOTES: {notes[:max_chars]}")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except BrokenPipeError:
        # Читатель закрыл пайп (`| head`, `| grep -q`) — это не ошибка.
        # Гасим traceback и глушим stdout, иначе интерпретатор напечатает
        # второй traceback уже при завершении.
        os.dup2(os.open(os.devnull, os.O_WRONLY), sys.stdout.fileno())
        sys.exit(0)
