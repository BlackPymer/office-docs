#!/bin/sh
"true" '''\'
# Python-скрипт: sh-трampoline ниже запускает его нужным интерпретатором.
# Порядок: .venv в корне репозитория -> .venv рядом со скриптом ->
# старый пользовательский venv -> python3.
D="$(dirname "$0")"
for P in "$D/../.venv/bin/python" "$D/.venv/bin/python" \
         "$D/../.venv/Scripts/python.exe" "$D/.venv/Scripts/python.exe" \
         "$HOME/.local/share/office-skills/venv/bin/python" \
         "$HOME/.local/share/office-skills/venv/Scripts/python.exe"; do
  [ -x "$P" ] && exec "$P" "$0" "$@"
done
exec python3 "$0" "$@"
''' + ''   # конкатенация -> Python не примет трampoline за docstring модуля


import os as _os
import sys as _sys

_sys.dont_write_bytecode = True  # не плодить __pycache__ рядом со скриптами скилла
_sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
from _common import ensure_runtime as _ensure_runtime  # noqa: E402

_ensure_runtime(("openpyxl",))  # на Windows shebang не работает: берём Python из venv

DOC = """Пересчёт формул .xlsx через LibreOffice и отчёт об ошибках.

Использование:
    xlsx_recalc.py <файл.xlsx> [таймаут_сек]

Файл перезаписывается на месте (LibreOffice считает формулы и сохраняет).
Печатает JSON: status, total_formulas, total_errors, error_summary.

Почему это обязательно: openpyxl пишет формулы как строки БЕЗ вычисленных
значений, поэтому pandas/openpyxl(data_only=True)/превьюеры видят None,
пока файл не пересчитан.

Предупреждения:
  - зелёный пересчёт доказывает, что формулы ВЫЧИСЛЯЮТСЯ, но не что верны;
    проверьте 2-3 ячейки вручную (xlsx_dump.py) перед финальной сдачей;
  - если книга ссылается на внешнюю книгу, пересчёт заменит ссылки на #NAME?.
"""
__doc__ = DOC


import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

sys.dont_write_bytecode = True  # не плодить __pycache__ рядом со скриптами скилла
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _common import (  # noqa: E402
    find_soffice,
    no_window_kwargs,
    soffice_env,
    user_installation_url,
)

from openpyxl import load_workbook  # noqa: E402

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>
"""

ERRORS = ["#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A"]


def stamp(p: Path):
    st = p.stat()
    return st.st_mtime_ns, st.st_size


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    path = Path(sys.argv[1]).expanduser().resolve()
    timeout = int(sys.argv[2]) if len(sys.argv) > 2 else 60
    if not path.is_file():
        print(json.dumps({"error": f"нет файла: {path}"}))
        return 1
    soffice = find_soffice()
    if not soffice:
        print(json.dumps({"error": "soffice не найден (см. SKILL.md)"}))
        return 1

    with tempfile.TemporaryDirectory(prefix="recalc_profile_") as prof:
        url = user_installation_url(prof)  # на Windows путь нужен как file:///C:/...
        # 1) bootstrap профиля, чтобы появились user/basic/Standard
        r = subprocess.run(
            [soffice, "--headless", "--terminate_after_init",
             f"-env:UserInstallation={url}"],
            env=soffice_env(), capture_output=True, text=True, timeout=timeout,
            **no_window_kwargs(),
        )
        macro_dir = Path(prof) / "user" / "basic" / "Standard"
        if not macro_dir.is_dir():
            print(json.dumps({"error": "LibreOffice не создал профиль — пересчёт не выполнен",
                              "detail": r.stderr.strip()}))
            return 1
        (macro_dir / "Module1.xba").write_text(MACRO)

        # 2) макрос поверх файла
        before = stamp(path)
        cmd = [soffice, "--headless", "--norestore", f"-env:UserInstallation={url}",
               "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?"
               "language=Basic&location=application",
               str(path)]
        started = time.monotonic()
        try:
            # Таймаут средствами Python: утилиты `timeout` на Windows нет.
            res = subprocess.run(cmd, env=soffice_env(), capture_output=True,
                                 text=True, timeout=timeout, **no_window_kwargs())
        except subprocess.TimeoutExpired:
            print(json.dumps({"error": f"LibreOffice завис (> {timeout}s) — формулы НЕ пересчитаны"}))
            return 1
        if res.returncode != 0:
            print(json.dumps({"error": f"soffice exit {res.returncode}",
                              "detail": (res.stderr or res.stdout).strip()[:500]}))
            return 1
        if stamp(path) == before:
            print(json.dumps({"error": "файл не перезаписан — пересчёт не произошёл "
                                       "(другой экземпляр LibreOffice?)"}, ensure_ascii=False))
            return 1

    # 3) отчёт об ошибках
    wb = load_workbook(path, data_only=True)
    total_formulas = 0
    wbf = load_workbook(path, data_only=False)
    for sname in wbf.sheetnames:
        for row in wbf[sname].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    total_formulas += 1
    summary: dict[str, list[str]] = {e: [] for e in ERRORS}
    total_errors = 0
    for sname in wb.sheetnames:
        ws = wb[sname]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value in summary:
                    total_errors += 1
                    if len(summary[c.value]) < 100:
                        summary[c.value].append(f"{sname}!{c.coordinate}")
    print(json.dumps({
        "status": "errors_found" if total_errors else "success",
        "total_formulas": total_formulas,
        "total_errors": total_errors,
        "error_summary": {k: v for k, v in summary.items() if v},
        "elapsed_s": round(time.monotonic() - started, 1),
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
