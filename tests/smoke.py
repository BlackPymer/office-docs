#!/bin/sh
"true" '''\'
D="$(dirname "$0")"
for P in "$D/../.venv/bin/python" "$D/.venv/bin/python" \
         "$D/../.venv/Scripts/python.exe" "$D/.venv/Scripts/python.exe" \
         "$HOME/.local/share/office-skills/venv/bin/python" \
         "$HOME/.local/share/office-skills/venv/Scripts/python.exe"; do
  [ -x "$P" ] && exec "$P" "$0" "$@"
done
exec python3 "$0" "$@"
''' + ''   # конкатенация -> Python не примет трampoline за docstring модуля
"""Самодостаточный smoke-тест office-docs.

Не зависит от чужих файлов: генерирует свои образцы (docx/xlsx/pptx/pdf) и
прогоняет по ним все скрипты скилла. Запускайте после установки:

    tests/smoke.py           # Linux/macOS
    py -3 tests\\smoke.py     # Windows

Ожидаемый итог — последняя строка `ИТОГ: PASS`. Любая FAIL-строка выше
означает, что что-то сломано или не установлено (см. сообщение).

Нужен установленный venv (install.py / install.sh). LibreOffice и poppler
проверяются, но их отсутствие не валит тест — только помечается как SKIP.
Ветки путей другой ОС проверяются через OFFICE_DOCS_FORCE_PLATFORM, поэтому
Windows-логика видна и на Linux-машине.
"""


import os as _os
import sys as _sys

_sys.dont_write_bytecode = True  # не плодить __pycache__ рядом со скриптами скилла
# smoke лежит в tests/, а _common — в scripts/: кладём оба каталога в sys.path
_TESTS_DIR = _os.path.dirname(_os.path.abspath(__file__))
_sys.path.insert(0, _os.path.join(_os.path.dirname(_TESTS_DIR), "scripts"))
from _common import ensure_runtime as _ensure_runtime  # noqa: E402

_ensure_runtime()  # на Windows shebang не работает: берём Python из venv

import itertools
import json
import os
import pathlib
import shutil
import subprocess
import sys
import tempfile
import textwrap
import zipfile

SKILL = pathlib.Path(__file__).resolve().parent.parent
SCRIPTS = SKILL / "scripts"
sys.path.insert(0, str(SCRIPTS))

from _common import find_pdftoppm, find_soffice, is_windows  # noqa: E402

passed: list[str] = []
failed: list[str] = []
skipped: list[str] = []


def child_env(**extra: str) -> dict:
    """Окружение для дочерних скриптов.

    Снимаем флаг ре-экзека, чтобы каждый скрипт сам искал интерпретатор, и
    фиксируем UTF-8: на Windows иначе кириллица в пайпах уезжает в cp1251.
    """
    env = os.environ.copy()
    env.pop("OFFICE_DOCS_BOOTSTRAPPED", None)
    env.setdefault("PYTHONIOENCODING", "utf-8")
    env.update(extra)
    return env


def script_cmd(script: str, *args: object) -> list[str]:
    """Как запускать скрипт скилла на этой платформе.

    POSIX — напрямую (заодно проверяем и shebang-трampoline, и бит выполнения).
    Windows — тем же интерпретатором, что и тест: shebang там не работает.
    """
    path = SCRIPTS / script
    tail = [str(a) for a in args]
    if is_windows():
        return [sys.executable, str(path), *tail]
    return [str(path), *tail]


def check(name: str, proc: subprocess.CompletedProcess, want=(0,), contains=None) -> None:
    ok = proc.returncode in want
    if ok and contains:
        ok = contains in proc.stdout
    if ok:
        passed.append(name)
    else:
        failed.append(name)
        print(f"  FAIL {name}: rc={proc.returncode}")
        if proc.stdout.strip():
            print(f"       stdout: {proc.stdout.strip()[:400]}")
        if proc.stderr.strip():
            print(f"       stderr: {proc.stderr.strip()[:400]}")


def expect(cond: bool, name: str) -> None:
    if cond:
        passed.append(name)
    else:
        failed.append(name)
        print(f"  FAIL {name}")


def run(name: str, script: str, *args: object, want=(0,), contains=None):
    p = subprocess.run(
        script_cmd(script, *args), capture_output=True, text=True, env=child_env()
    )
    check(name, p, want, contains)
    return p


def shell(name: str, *args: object, want=(0,), contains=None, env: dict | None = None):
    p = subprocess.run(
        [str(a) for a in args], capture_output=True, text=True, env=env or child_env()
    )
    check(name, p, want, contains)
    return p


# --------------------------------------------------------------------------
# Проверка путей другой ОС: гоняем _common в отдельном процессе с подменой
# платформы (OFFICE_DOCS_FORCE_PLATFORM) — так Windows-ветки тестируются и здесь.
# --------------------------------------------------------------------------
PROBE = textwrap.dedent(
    r'''
    import json, os, sys
    sys.path.insert(0, os.environ["SKILL_SCRIPTS"])
    import _common
    print(json.dumps({
        "windows": _common.is_windows(),
        "uri": _common.user_installation_url(os.environ["PROBE_PATH"]),
        "soffice": _common.soffice_candidates(),
        "venv": _common.venv_python_candidates(),
    }, ensure_ascii=False))
    '''
)


def probe(**env_extra: str) -> dict:
    p = subprocess.run(
        [sys.executable, "-c", PROBE],
        capture_output=True,
        text=True,
        env=child_env(SKILL_SCRIPTS=str(SCRIPTS), **env_extra),
    )
    if p.returncode != 0:
        failed.append("проба _common")
        print(f"  FAIL проба _common: {p.stderr.strip()[:300]}")
        return {}
    return json.loads(p.stdout)


def test_platform_paths() -> None:
    print("== Пути Windows и POSIX (подмена платформы) ==")
    win_path = r"C:\Users\me\AppData\Local\Temp\lo_profile"
    win = probe(
        OFFICE_DOCS_FORCE_PLATFORM="windows",
        PROBE_PATH=win_path,
        ProgramFiles=r"C:\Program Files",
        ProgramW6432=r"C:\Program Files",
        LOCALAPPDATA=r"C:\Users\me\AppData\Local",
        SOFFICE=r"C:\Tools\LibreOffice\program\soffice.exe",
    )
    if win:
        expect(win["windows"] is True, "платформа: Windows-ветка включается")
        expect(
            win["uri"] == "file:///C:/Users/me/AppData/Local/Temp/lo_profile",
            "Windows: -env:UserInstallation как file:///C:/...",
        )
        expect(
            win["soffice"][:1] == [r"C:\Tools\LibreOffice\program\soffice.exe"],
            "Windows: SOFFICE=... имеет высший приоритет",
        )
        expect(
            any(p.replace("\\", "/").endswith("LibreOffice/program/soffice.exe")
                for p in win["soffice"]),
            "Windows: находится установка в Program Files",
        )
        expect(
            any(p.replace("\\", "/").endswith(".venv/Scripts/python.exe") for p in win["venv"]),
            "Windows: venv-интерпретатор ищется в .venv/Scripts/python.exe",
        )
        expect(
            any(p.replace("\\", "/").endswith("scoop/apps/libreoffice/current/program/soffice.exe")
                for p in win["soffice"]),
            "Windows: находится scoop-установка",
        )

    posix = probe(OFFICE_DOCS_FORCE_PLATFORM="posix", PROBE_PATH="/tmp/lo_profile")
    if posix:
        expect(posix["windows"] is False, "платформа: POSIX-ветка включается")
        expect(posix["uri"] == "file:///tmp/lo_profile", "POSIX: -env:UserInstallation как file:///tmp/...")
        expect(
            any(p.replace("\\", "/").endswith(".venv/bin/python") for p in posix["venv"]),
            "POSIX: venv-интерпретатор ищется в .venv/bin/python",
        )


def test_broken_pipe() -> None:
    """Скрипт, у которого закрыли пайп (`| head`), не должен падать с traceback."""
    print("== Пайпы ==")
    proc = subprocess.Popen(
        script_cmd("pptx_outline.py", "deck.pptx"),
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        env=child_env(),
    )
    head = "".join(itertools.islice(proc.stdout, 3))
    proc.stdout.close()  # имитируем head, закрывший пайп
    err = proc.stderr.read()
    proc.wait()
    expect("Traceback" not in err, "обрыв пайпа не даёт traceback")
    expect(bool(head.strip()), "pptx_outline что-то пишет в пайп")


def main() -> int:
    work = pathlib.Path(tempfile.mkdtemp(prefix="office_smoke_"))
    os.chdir(work)
    print(f"Рабочий каталог: {work}")
    print(f"Платформа: {'Windows' if is_windows() else os.uname().sysname if hasattr(os, 'uname') else os.name}")

    # --- образцы -------------------------------------------------------
    import docx, openpyxl, pymupdf
    from pptx import Presentation
    from pptx.util import Inches

    d = docx.Document()
    d.add_heading("Smoke test", 0)
    d.add_paragraph("Абзац с кириллицей для проверки шрифтов.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Имя", "Значение"
    t.cell(1, 0).text, t.cell(1, 1).text = "x", "42"
    d.save("in.docx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["A2"], ws["A3"], ws["B1"] = "x", 3, 4, "=A2+A3"
    wb.save("book.xlsx")

    prs = Presentation()
    for i in (1, 2):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        box = s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        box.text_frame.text = f"Слайд {i}"
    prs.save("deck.pptx")

    pdf = pymupdf.open()
    for i in (1, 2):
        page = pdf.new_page()
        page.insert_text((72, 72), f"Page {i}")
    pdf.save("doc.pdf")
    pdf.close()

    # --- LibreOffice: нужен нескольким проверкам, определяем один раз -------
    soffice = find_soffice()
    pdftoppm = find_pdftoppm()
    if soffice:
        print(f"(LibreOffice: {soffice})")

    test_platform_paths()

    # --- Excel ---------------------------------------------------------
    print("\n== Excel ==")
    run("xlsx_dump", "xlsx_dump.py", "book.xlsx")
    run("xlsx_dump --formulas", "xlsx_dump.py", "book.xlsx", "--formulas",
        contains="=A2+A3")
    if soffice:
        run("xlsx_recalc", "xlsx_recalc.py", "book.xlsx", contains='"total_errors": 0')
    else:
        skipped.append("xlsx_recalc — нужен LibreOffice (пересчёт формул)")

    # --- PowerPoint ----------------------------------------------------
    print("== PowerPoint ==")
    run("pptx_outline", "pptx_outline.py", "deck.pptx", contains="Слайд 1")
    run("pptx_dup_slide", "pptx_dup_slide.py", "deck.pptx",
        "--source", "1", "--after", "1", "-o", "dup.pptx")
    run("дубль дал 3 слайда", "pptx_outline.py", "dup.pptx", contains="3 слайд")

    # --- PDF -----------------------------------------------------------
    print("== PDF ==")
    run("pdf_tool info", "pdf_tool.py", "info", "doc.pdf", contains="страниц: 2")
    run("pdf_tool text", "pdf_tool.py", "text", "doc.pdf", contains="Page 1")
    run("pdf_tool rotate", "pdf_tool.py", "rotate", "doc.pdf",
        "-o", "rot.pdf", "--deg", "90")
    run("pdf_tool merge", "pdf_tool.py", "merge", "-o", "m.pdf", "doc.pdf", "doc.pdf")
    run("pdf_tool split", "pdf_tool.py", "split", "doc.pdf",
        "-o", "s.pdf", "--pages", "1-1")

    # --- Word ----------------------------------------------------------
    print("== Word ==")
    run("merge_runs", "merge_runs.py", "in.docx", "-o", "out.docx")

    # --- свой код интерпретатором из venv ------------------------------
    print("== Python из venv (py.py) ==")
    run("py.py -c", "py.py", "-c",
        "import docx, openpyxl, pymupdf; print('PYOK')", contains="PYOK")

    # --- ZIP-пакет OOXML (замена unzip/zip) ----------------------------
    print("== OOXML-пакет ==")
    run("ooxml_zip list", "ooxml_zip.py", "list", "in.docx", contains="word/document.xml")
    run("ooxml_zip cat", "ooxml_zip.py", "cat", "in.docx", "word/document.xml",
        contains="<w:document")
    run("ooxml_zip unpack", "ooxml_zip.py", "unpack", "in.docx", "unpacked")
    unpacked = pathlib.Path("unpacked")
    expect((unpacked / "word" / "document.xml").is_file(),
           "ooxml_zip unpack разложил части по каталогам")
    # мусор рядом с пакетом не должен попасть в архив
    if unpacked.is_dir():
        (unpacked / ".DS_Store").write_text("junk")
        run("ooxml_zip pack", "ooxml_zip.py", "pack", "unpacked", "repacked.docx")
        listing = run("ooxml_zip list repacked", "ooxml_zip.py", "list", "repacked.docx")
        expect(".DS_Store" not in listing.stdout, "ooxml_zip pack не тащит мусор")
        expect("word/document.xml" in listing.stdout, "ooxml_zip pack сохранил части пакета")
        if pathlib.Path("repacked.docx").is_file():
            with zipfile.ZipFile("repacked.docx") as z:
                expect(z.read("word/document.xml").startswith(b"<?xml"),
                       "ooxml_zip pack: XML внутри читается как XML")
            # собранный пакет должен оставаться OOXML с точки зрения наших же проверок
            run("validate repacked", "validate.py", "repacked.docx", contains="ИТОГ: PASS")

    # --- LibreOffice: конвертация, валидация, рендер --------------------
    if not soffice:
        skipped.append("LibreOffice не найден: office_convert, render_check, validate --soffice")
        print("  SKIP: LibreOffice не найден, конвертация и рендер пропущены")
    else:
        run("office_convert docx->pdf", "office_convert.py", "in.docx", "c.pdf")
        run("office_convert docx->md", "office_convert.py", "in.docx", "c.md")
        run("validate --soffice", "validate.py", "in.docx", "--soffice", contains="ИТОГ: PASS")

        print("== render_check.py ==")
        run("render_check.py in.docx", "render_check.py", "in.docx")
        imgs = sorted(pathlib.Path(".").glob("in-*.jpg"))
        expect(bool(imgs), "render_check дал картинки страниц")
        expect(all(i.read_bytes()[:2] == b"\xff\xd8" for i in imgs),
               "render_check: страницы — настоящие JPEG")
        run("render_check.py из PDF", "render_check.py", "doc.pdf", "pdfpages", "72")
        pdf_pages = sorted(pathlib.Path(".").glob("pdfpages-*.jpg"))
        expect([p.name for p in pdf_pages] == ["pdfpages-01.jpg", "pdfpages-02.jpg"],
               "render_check: PDF на входе даёт обе страницы постранично")

        if not is_windows():
            shell("render_check.sh (обёртка)", "bash", str(SCRIPTS / "render_check.sh"),
                  "in.docx", "shim")
            expect(sorted(pathlib.Path(".").glob("shim-*.jpg")),
                   "render_check.sh вызывает render_check.py")

        if pdftoppm:
            # запасной путь без PyMuPDF тоже должен работать
            code = textwrap.dedent(
                """
                import os, pathlib, sys, tempfile
                sys.path.insert(0, os.environ["SKILL_SCRIPTS"])
                import render_check
                with tempfile.TemporaryDirectory() as tmp:
                    os.chdir(tmp)
                    n = render_check.rasterize_poppler(
                        os.environ["PDFTOPPM"], pathlib.Path(os.environ["SAMPLE_PDF"]), "pop", 60)
                    print(n, sorted(p.name for p in pathlib.Path(tmp).glob("*.jpg")))
                """
            )
            p = subprocess.run(
                [sys.executable, "-c", code],
                capture_output=True,
                text=True,
                env=child_env(
                    SKILL_SCRIPTS=str(SCRIPTS),
                    PDFTOPPM=pdftoppm,
                    SAMPLE_PDF=str(pathlib.Path("doc.pdf").resolve()),
                ),
            )
            check("render_check через pdftoppm", p, contains="pop-01.jpg")
        else:
            skipped.append("render_check через pdftoppm — нет pdftoppm (не обязателен)")

    test_broken_pipe()

    print("== Без аргументов печатают usage и код 2 ==")
    for s in ("validate", "office_convert", "pdf_tool", "xlsx_dump", "xlsx_recalc",
              "pptx_outline", "merge_runs", "render_check", "ooxml_zip", "py"):
        run(f"{s}.py без аргументов", f"{s}.py", want=(2,))

    print()
    for s in skipped:
        print(f"SKIP: {s}")
    print(f"ИТОГ: PASS={len(passed)}  FAIL={len(failed)}  SKIP={len(skipped)}")
    shutil.rmtree(work, ignore_errors=True)
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
